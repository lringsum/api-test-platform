from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import time
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import requests
from requests.exceptions import RequestException
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment


SOURCE_XLSX = Path(r"C:\Users\Administrator\WorkBuddy\20260702170507\apk_size_baseline_20260706.xlsx")
REPORT_DIR = Path(r"C:\Users\Administrator\WorkBuddy\20260702170507\apk_ui_check_20260706")
REPORT_XLSX = REPORT_DIR / "apk_ui_report_20260706.xlsx"
PROGRESS_LOG = REPORT_DIR / "apk_ui_progress_20260706.log"
APK_DIR = REPORT_DIR / "apks"
SCREENSHOT_DIR = REPORT_DIR / "screenshots"
PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _resolve_adb_path() -> Path:
    candidates = [
        os.environ.get("ANDROID_UI_ADB_PATH", "").strip(),
        r"C:\Program Files\Netease\MuMu\nx_main\adb.exe",
    ]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    adb_in_path = shutil.which("adb")
    if adb_in_path:
        return Path(adb_in_path)
    raise FileNotFoundError("未找到 adb，请先安装 MuMu，或设置 ANDROID_UI_ADB_PATH。")


def _resolve_aapt_path() -> Path:
    candidates = [
        os.environ.get("ANDROID_UI_AAPT_PATH", "").strip(),
        str(PROJECT_ROOT / "tools" / "android" / "aapt.exe"),
    ]
    android_homes = [
        os.environ.get("ANDROID_HOME") or os.environ.get("ANDROID_SDK_ROOT"),
        str(Path(os.environ.get("LOCALAPPDATA", "")) / "Android" / "Sdk"),
    ]
    for android_home in android_homes:
        if not android_home:
            continue
        build_tools_root = Path(android_home) / "build-tools"
        if not build_tools_root.is_dir():
            continue
        for version_dir in sorted(build_tools_root.iterdir(), reverse=True):
            candidates.append(str(version_dir / "aapt.exe"))
            candidates.append(str(version_dir / "aapt"))
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return Path(candidate)
    raise FileNotFoundError("未找到 aapt，请先安装 Android SDK build-tools，或设置 ANDROID_UI_AAPT_PATH。")


ADB = _resolve_adb_path()
AAPT = _resolve_aapt_path()
SERIAL = os.environ.get("ANDROID_UI_SERIAL", "127.0.0.1:16384").strip() or "127.0.0.1:16384"


HEADERS = [
    "\u5305\u6807\u8bc6",
    "\u5305\u540d\u79f0",
    "\u5305\u94fe\u63a5",
    "\u5b89\u88c5\u7ed3\u679c",
    "\u542f\u52a8\u7ed3\u679c",
    "\u8fd0\u884cPID",
    "\u7126\u70b9\u7a97\u53e3",
    "\u95ea\u9000\u5224\u65ad",
    "\u754c\u9762\u622a\u56fe",
    "\u5907\u6ce8",
    "\u68c0\u6d4b\u65f6\u95f4",
]


@dataclass
class InputItem:
    pkg_flag: str
    url: str


@dataclass
class SourceRow:
    pkg_flag: str
    pkg_name: str
    url: str


@dataclass
class RunResult:
    pkg_flag: str
    pkg_name: str
    url: str
    install_result: str
    launch_result: str
    pid: str
    focus: str
    crash_judgement: str
    note: str
    checked_at: str
    screenshot_path: Optional[Path]


def append_progress(log_path: Path, message: str) -> None:
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with log_path.open("a", encoding="utf-8") as fh:
        fh.write(f"[{timestamp}] {message}\n")


def run_cmd(args: List[str], timeout: int = 120, check: bool = False) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        args,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=timeout,
        check=check,
    )


def load_input_items(path: Path) -> List[InputItem]:
    items: List[InputItem] = []
    for raw in path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if not line or line.startswith("#"):
            continue
        parts = line.split("\t")
        if len(parts) != 2:
            raise ValueError(f"Invalid line in input file: {raw!r}")
        items.append(InputItem(pkg_flag=parts[0].strip(), url=parts[1].strip()))
    return items


def load_source_map(path: Path) -> Dict[str, SourceRow]:
    wb = load_workbook(path, data_only=True)
    ws = wb["baseline_20260702"]
    result: Dict[str, SourceRow] = {}
    for row in range(2, ws.max_row + 1):
        pkg_flag = str(ws.cell(row, 2).value or "").strip()
        pkg_name = str(ws.cell(row, 3).value or "").strip()
        url = str(ws.cell(row, 5).value or ws.cell(row, 4).value or "").strip()
        if pkg_flag:
            result[pkg_flag] = SourceRow(pkg_flag=pkg_flag, pkg_name=pkg_name, url=url)
    return result


def ensure_dirs() -> None:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    APK_DIR.mkdir(parents=True, exist_ok=True)
    SCREENSHOT_DIR.mkdir(parents=True, exist_ok=True)


def save_workbook_with_retry(wb, path: Path, retries: int = 5, delay_seconds: int = 3) -> None:
    last_error: Optional[Exception] = None
    for _ in range(retries):
        try:
            temp_path = path.with_name(f"{path.stem}.__tmp__{int(time.time() * 1000)}{path.suffix}")
            wb.save(temp_path)
            os.replace(temp_path, path)
            return
        except PermissionError as exc:
            last_error = exc
            time.sleep(delay_seconds)
        finally:
            try:
                if "temp_path" in locals() and Path(temp_path).exists():
                    Path(temp_path).unlink()
            except OSError:
                pass
    if last_error:
        raise last_error


def rebuild_report_if_needed(report_path: Path, source_map: Dict[str, SourceRow]) -> None:
    if not report_path.exists():
        create_empty_report(report_path)
        return

    old = load_workbook(report_path)
    old_ws = old.active
    should_rebuild = False
    header_values = [old_ws.cell(1, idx + 1).value for idx in range(len(HEADERS))]
    if header_values != HEADERS:
        should_rebuild = True
    else:
        for row in range(2, old_ws.max_row + 1):
            value = old_ws.cell(row, 2).value
            if isinstance(value, str) and "?" in value:
                should_rebuild = True
                break

    if not should_rebuild:
        return

    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = "ui_check_20260706"
    write_headers(new_ws)

    for row in range(2, old_ws.max_row + 1):
        pkg_flag = str(old_ws.cell(row, 1).value or "").strip()
        if not pkg_flag:
            continue
        src = source_map.get(pkg_flag)
        pkg_name = src.pkg_name if src else str(old_ws.cell(row, 2).value or "").strip()
        url = src.url if src else str(old_ws.cell(row, 3).value or "").strip()
        values = [
            pkg_flag,
            pkg_name,
            url,
            old_ws.cell(row, 5).value or "",
            old_ws.cell(row, 7).value or "",
            old_ws.cell(row, 8).value or "",
            old_ws.cell(row, 9).value or "",
            "\u672a\u5224\u5b9a",
            "",
            old_ws.cell(row, 6).value or "",
            datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        ]
        dst_row = append_row(new_ws, values)
        old_path = str(old_ws.cell(row, 10).value or "").strip()
        if old_path and Path(old_path).exists():
            embed_image(new_ws, dst_row, Path(old_path))

    save_workbook_with_retry(new_wb, report_path)


def create_empty_report(report_path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ui_check_20260706"
    write_headers(ws)
    save_workbook_with_retry(wb, report_path)


def write_headers(ws) -> None:
    for idx, header in enumerate(HEADERS, start=1):
        ws.cell(1, idx, header)
    ws.freeze_panes = "A2"
    widths = {
        "A": 14,
        "B": 28,
        "C": 70,
        "D": 18,
        "E": 28,
        "F": 12,
        "G": 50,
        "H": 14,
        "I": 36,
        "J": 22,
        "K": 20,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def append_row(ws, values: List[str]) -> int:
    row = ws.max_row + 1
    for idx, value in enumerate(values, start=1):
        ws.cell(row, idx, value)
        ws.cell(row, idx).alignment = Alignment(vertical="top", wrap_text=True)
    ws.row_dimensions[row].height = 160
    return row


def embed_image(ws, row: int, image_path: Path) -> None:
    img = XLImage(str(image_path))
    img.width = 280
    img.height = 158
    ws.add_image(img, f"I{row}")


def read_report_rows(report_path: Path) -> List[List[str]]:
    if not report_path.exists():
        return []
    wb = load_workbook(report_path, data_only=True)
    ws = wb.active
    rows: List[List[str]] = []
    for row in range(2, ws.max_row + 1):
        pkg_flag = str(ws.cell(row, 1).value or "").strip()
        if not pkg_flag:
            continue
        values = [ws.cell(row, col).value or "" for col in range(1, len(HEADERS) + 1)]
        rows.append(values)
    return rows


def write_report_rows(report_path: Path, rows: List[List[str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "ui_check_20260706"
    write_headers(ws)
    for values in rows:
        row_index = append_row(ws, values)
        pkg_flag = str(values[0] or "").strip()
        image_path = SCREENSHOT_DIR / f"{pkg_flag}.png"
        if image_path.exists():
            embed_image(ws, row_index, image_path)
    save_workbook_with_retry(wb, report_path)


def row_lookup(ws) -> Dict[str, int]:
    mapping: Dict[str, int] = {}
    for row in range(2, ws.max_row + 1):
        key = str(ws.cell(row, 1).value or "").strip()
        if key:
            mapping[key] = row
    return mapping


def download_apk(url: str, dst: Path, retries: int = 3) -> None:
    if dst.exists() and dst.stat().st_size > 0:
        return
    last_error: Optional[Exception] = None
    for _ in range(retries):
        try:
            if dst.exists():
                dst.unlink()
            with requests.get(url, stream=True, timeout=90) as resp:
                resp.raise_for_status()
                with dst.open("wb") as fh:
                    for chunk in resp.iter_content(chunk_size=1024 * 1024):
                        if chunk:
                            fh.write(chunk)
            if dst.exists() and dst.stat().st_size > 0:
                return
        except RequestException as exc:
            last_error = exc
            time.sleep(3)
        except OSError as exc:
            last_error = exc
            time.sleep(1)
    if last_error:
        raise last_error


def parse_badging(apk_path: Path) -> Tuple[str, Optional[str]]:
    badging = run_cmd([str(AAPT), "dump", "badging", str(apk_path)], timeout=180, check=True).stdout
    package_name = ""
    launchable = None
    for line in badging.splitlines():
        if line.startswith("package:"):
            marker = "name='"
            package_name = line.split(marker, 1)[1].split("'", 1)[0]
        elif line.startswith("launchable-activity:"):
            marker = "name='"
            launchable = line.split(marker, 1)[1].split("'", 1)[0]
    if not package_name:
        raise RuntimeError(f"Unable to resolve package name for {apk_path}")
    return package_name, launchable


def adb(args: List[str], timeout: int = 120) -> subprocess.CompletedProcess[str]:
    return run_cmd([str(ADB), "-s", SERIAL, *args], timeout=timeout)


def ensure_device() -> None:
    out = adb(["get-state"], timeout=30)
    if out.returncode != 0 or "device" not in out.stdout:
        raise RuntimeError("ADB device is not ready")


def capture_screenshot(dst: Path) -> None:
    shot = subprocess.run(
        [str(ADB), "-s", SERIAL, "exec-out", "screencap", "-p"],
        capture_output=True,
        timeout=120,
        check=True,
    )
    dst.write_bytes(shot.stdout)


def launch_app(package_name: str, launchable: Optional[str]) -> str:
    if launchable:
        cp = adb(["shell", "am", "start", "-n", f"{package_name}/{launchable}"], timeout=120)
    else:
        cp = adb(
            ["shell", "monkey", "-p", package_name, "-c", "android.intent.category.LAUNCHER", "1"],
            timeout=120,
        )
    return summarize_output(cp)


def summarize_output(cp: subprocess.CompletedProcess[str]) -> str:
    text = (cp.stdout + "\n" + cp.stderr).strip()
    return " ".join(line.strip() for line in text.splitlines() if line.strip())[:500]


def inspect_running(package_name: str) -> Tuple[str, str, str]:
    pid = adb(["shell", "pidof", package_name], timeout=60).stdout.strip()
    focus_dump = adb(["shell", "dumpsys", "window", "windows"], timeout=180).stdout
    focus_lines = [
        line.strip()
        for line in focus_dump.splitlines()
        if "mCurrentFocus" in line or "mFocusedApp" in line
    ]
    focus = "\n".join(focus_lines)[:1000]
    if pid and package_name in focus:
        judgement = "\u672a\u53d1\u73b0\u95ea\u9000"
    elif pid:
        judgement = "\u8fd0\u884c\u4e2d\uff0c\u4f46\u7126\u70b9\u4e0d\u5728\u6e38\u620f"
    else:
        judgement = "\u7591\u4f3c\u95ea\u9000\u6216\u9000\u51fa"
    return pid, focus, judgement


def uninstall_app(package_name: str) -> None:
    adb(["uninstall", package_name], timeout=240)


def process_item(
    item: InputItem,
    source_map: Dict[str, SourceRow],
    wait_seconds: int,
    install_timeout: int,
    progress_log: Path,
    index: int,
    total: int,
) -> RunResult:
    src = source_map.get(item.pkg_flag)
    pkg_name = src.pkg_name if src else item.pkg_flag
    url = item.url or (src.url if src else "")
    checked_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    apk_path = APK_DIR / f"{item.pkg_flag}.apk"
    screenshot_path = SCREENSHOT_DIR / f"{item.pkg_flag}.png"
    note = ""
    package_name = ""
    append_progress(progress_log, f"START {index}/{total} {item.pkg_flag} {pkg_name} {url}")

    try:
        ensure_device()
        download_apk(url, apk_path)
        package_name, launchable = parse_badging(apk_path)
        uninstall_app(package_name)
        install_cp = adb(["install", "-r", "-g", str(apk_path)], timeout=install_timeout)
        install_result = summarize_output(install_cp)
        if "Success" not in install_result:
            note = "\u5b89\u88c5\u672a\u6210\u529f"
            result = RunResult(
                item.pkg_flag,
                pkg_name,
                url,
                install_result,
                "",
                "",
                "",
                "\u5b89\u88c5\u5931\u8d25",
                note,
                checked_at,
                None,
            )
            append_progress(progress_log, f"DONE {index}/{total} {item.pkg_flag} 安装失败 {note}")
            return result

        launch_result = launch_app(package_name, launchable)
        time.sleep(wait_seconds)
        pid, focus, crash_judgement = inspect_running(package_name)
        capture_screenshot(screenshot_path)
        result = RunResult(
            item.pkg_flag,
            pkg_name,
            url,
            install_result,
            launch_result,
            pid,
            focus,
            crash_judgement,
            note,
            checked_at,
            screenshot_path,
        )
        append_progress(progress_log, f"DONE {index}/{total} {item.pkg_flag} {crash_judgement}")
        return result
    except Exception as exc:
        note = f"{type(exc).__name__}: {exc}"
        result = RunResult(
            item.pkg_flag,
            pkg_name,
            url,
            "",
            "",
            "",
            "",
            "\u6267\u884c\u5f02\u5e38",
            note,
            checked_at,
            None,
        )
        append_progress(progress_log, f"DONE {index}/{total} {item.pkg_flag} 执行异常 {note}")
        return result
    finally:
        if package_name:
            adb(["shell", "am", "force-stop", package_name], timeout=120)
            uninstall_app(package_name)


def upsert_result(report_path: Path, result: RunResult) -> None:
    existing_rows = read_report_rows(report_path)
    mapping: Dict[str, int] = {}
    for idx, row in enumerate(existing_rows):
        key = str(row[0] or "").strip()
        if key:
            mapping[key] = idx
    values = [
        result.pkg_flag,
        result.pkg_name,
        result.url,
        result.install_result,
        result.launch_result,
        result.pid,
        result.focus,
        result.crash_judgement,
        "",
        result.note,
        result.checked_at,
    ]

    if result.pkg_flag in mapping:
        existing_rows[mapping[result.pkg_flag]] = values
    else:
        existing_rows.append(values)

    write_report_rows(report_path, existing_rows)


def parse_args(argv: Optional[Iterable[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True, type=Path)
    parser.add_argument("--wait-seconds", type=int, default=35)
    parser.add_argument("--install-timeout", type=int, default=900)
    parser.add_argument("--report", type=Path, default=REPORT_XLSX)
    parser.add_argument("--report-template", type=Path, default=REPORT_XLSX)
    parser.add_argument("--progress-log", type=Path, default=PROGRESS_LOG)
    return parser.parse_args(argv)


def main(argv: Optional[Iterable[str]] = None) -> int:
    args = parse_args(argv)
    ensure_dirs()
    source_map = load_source_map(SOURCE_XLSX)
    items = load_input_items(args.input)
    if args.report != args.report_template and not args.report.exists() and args.report_template.exists():
        template_rows = read_report_rows(args.report_template)
        write_report_rows(args.report, template_rows)

    rebuild_report_if_needed(args.report, source_map)
    if not args.report.exists():
        create_empty_report(args.report)

    append_progress(args.progress_log, f"RUN START total={len(items)} report={args.report}")
    for idx, item in enumerate(items, start=1):
        result = process_item(
            item,
            source_map,
            args.wait_seconds,
            args.install_timeout,
            args.progress_log,
            idx,
            len(items),
        )
        upsert_result(args.report, result)
        print(f"{result.pkg_flag}\t{result.crash_judgement}\t{result.note or '-'}", flush=True)

    append_progress(args.progress_log, f"RUN END total={len(items)} report={args.report}")
    print(str(args.report))
    return 0


if __name__ == "__main__":
    sys.exit(main())
